"""Тесты отчётов преподавателя: посещаемость и должники (просмотр и выгрузка в Excel)."""
import io
from datetime import date, timedelta
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Attendance, Defense
from core.tests.helpers import make_lab_work, make_report, make_student, make_subgroup, make_teacher

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _sheet_rows(response):
    """Читает скачанный xlsx-файл и возвращает его строки списком."""
    workbook = load_workbook(io.BytesIO(response.content))
    return [list(row) for row in workbook.active.iter_rows(values_only=True)]


class ReportsTestCase(TestCase):
    def setUp(self):
        self.teacher = make_teacher()
        self.subgroup = make_subgroup(self.teacher)
        self.student = make_student(self.subgroup)
        self.client.force_login(self.teacher.profile.user)


class AttendanceReportTests(ReportsTestCase):
    def setUp(self):
        super().setUp()
        Attendance.objects.create(student=self.student, lesson_date=date(2024, 3, 1), status='П')
        Attendance.objects.create(student=self.student, lesson_date=date(2024, 3, 8), status='Н')
        Attendance.objects.create(student=self.student, lesson_date=date(2024, 3, 15), status='Б')


    def test_subgroup_filter(self):
        other = make_student(make_subgroup(self.teacher))
        response = self.client.get(reverse('attendance_report'), {'subgroup': self.subgroup.pk})
        self.assertEqual([r['student'] for r in response.context['rows']], [self.student])
        self.assertNotIn(other, [r['student'] for r in response.context['rows']])

    def test_xlsx_export(self):
        response = self.client.get(reverse('attendance_report'), {'export': 'xlsx'})
        self.assertEqual(response['Content-Type'], XLSX)
        self.assertIn('attendance_report.xlsx', response['Content-Disposition'])
        rows = _sheet_rows(response)
        self.assertEqual(rows[0][0], 'Отчёт по посещаемости')
        self.assertEqual(rows[2][-2:], [1, 2])


class DebtorsReportTests(ReportsTestCase):
    def setUp(self):
        super().setUp()
        self.lab_work = make_lab_work()  # дедлайны в 2024 году — уже прошли


    def test_xlsx_export(self):
        response = self.client.get(reverse('debtors_report'), {'export': 'xlsx', 'subgroup': self.subgroup.pk})
        self.assertEqual(response['Content-Type'], XLSX)
        rows = _sheet_rows(response)
        self.assertEqual(rows[1][:2], ['Студент', 'Подгруппа'])
        self.assertEqual(len(rows), 4)  # заголовок, шапка и два долга